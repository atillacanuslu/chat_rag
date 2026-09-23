"""Turn an eval CSV into something readable.

    python eval/report.py eval/results/baseline.csv
        prints a compact table plus a summary in the terminal

    python eval/report.py eval/results/baseline.csv --xlsx eval/results/baseline.xlsx
        also writes a formatted workbook: Summary, Questions, Scores

    python eval/report.py eval/results/baseline.csv eval/results/after.csv --xlsx cmp.xlsx
        compares two runs side by side (before and after)
"""

import argparse
import csv
import json
import os
import statistics


def read(path):
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def summarize(rows):
    scored = [r for r in rows if r.get("scored") == "1" and not r.get("error")]
    judged = [r for r in rows if r.get("correct") in ("0", "1")]
    times = [num(r["seconds"]) for r in rows if not r.get("error")]
    times = [t for t in times if t is not None]
    ptok = [num(r.get("prompt_tokens")) for r in rows]
    ptok = [t for t in ptok if t is not None]
    nan = [r for r in rows if num(r.get("rerank_nan_count"))]
    zeros = [num(r.get("bm25_zero_count")) for r in rows]
    zeros = [z for z in zeros if z is not None]

    def spread(key):
        vals = [num(r.get(key)) for r in rows]
        vals = [v for v in vals if v is not None]
        return statistics.median(vals) if vals else None

    s = {
        "questions": len(rows),
        "scored": len(scored),
        "hit_rate": (sum(1 for r in scored if r["hit_final"] == "1") / len(scored) * 100) if scored else None,
        "mrr": (sum(num(r["rr"]) or 0 for r in scored) / len(scored)) if scored else None,
        "accuracy": (sum(1 for r in judged if r["correct"] == "1") / len(judged) * 100) if judged else None,
        "crash_rate": sum(1 for r in rows if r.get("error")) / len(rows) * 100,
        "nan_queries": (len(nan) / len(rows) * 100) if rows else None,
        "median_latency": statistics.median(times) if times else None,
        "p90_latency": sorted(times)[int(len(times) * 0.9) - 1] if times else None,
        "median_prompt_tokens": statistics.median(ptok) if ptok else None,
        "llm_calls": rows[0].get("llm_calls", "") if rows else "",
        "vector_spread": spread("vec_spread"),
        "bm25_spread": spread("bm25_spread"),
        "hybrid_spread": spread("hybrid_spread"),
        "bm25_zero_avg": (sum(zeros) / len(zeros)) if zeros else None,
    }
    return s


LABELS = [
    ("questions", "questions", "{:.0f}"),
    ("hit_rate", "hit rate @5 (%)", "{:.1f}"),
    ("mrr", "MRR", "{:.3f}"),
    ("accuracy", "answer accuracy (%)", "{:.1f}"),
    ("crash_rate", "crash rate (%)", "{:.1f}"),
    ("nan_queries", "queries with NaN rerank (%)", "{:.1f}"),
    ("median_latency", "median latency (s)", "{:.1f}"),
    ("p90_latency", "p90 latency (s)", "{:.1f}"),
    ("median_prompt_tokens", "median input tokens", "{:.0f}"),
    ("llm_calls", "LLM calls per question", "{}"),
    ("vector_spread", "vector score spread", "{:.4f}"),
    ("bm25_spread", "bm25 score spread", "{:.4f}"),
    ("hybrid_spread", "hybrid score spread", "{:.4f}"),
    ("bm25_zero_avg", "bm25 zero-score results (avg)", "{:.1f}"),
]


def fmt(val, pattern):
    if val is None or val == "":
        return "-"
    try:
        return pattern.format(val)
    except (ValueError, TypeError):
        return str(val)


def print_terminal(paths, runs, summaries):
    w = 30
    print("\nSUMMARY")
    print("-" * (w + 14 * len(runs)))
    header = "metric".ljust(w) + "".join(os.path.basename(p)[:12].rjust(14) for p in paths)
    print(header)
    print("-" * (w + 14 * len(runs)))
    for key, label, pattern in LABELS:
        line = label.ljust(w) + "".join(fmt(s[key], pattern).rjust(14) for s in summaries)
        print(line)
    print("-" * (w + 14 * len(runs)))

    rows = runs[0]
    print("\nPER QUESTION  (%s)" % os.path.basename(paths[0]))
    print("-" * 96)
    print("%-5s %-4s %-5s %-6s %-40s %-14s %s" % (
        "id", "hit", "rank", "judge", "question", "expected", "answer"))
    print("-" * 150)
    for r in rows:
        if r.get("error"):
            status, rank, judge = "ERR", "-", "-"
            tail = r["error"][:60]
        else:
            status = "hit" if r.get("hit_final") == "1" else ("-" if r.get("hit_final") == "" else "MISS")
            rank = r.get("rank") or "-"
            judge = {"1": "ok", "0": "wrong", "": "?"}.get(r.get("correct", ""), "-")
            tail = (r.get("answer") or "")[:60]
        print("%-5s %-4s %-5s %-6s %-40s %-14s %s" % (
            r["id"], status, rank, judge,
            (r.get("question") or "")[:40], (r.get("expected") or "")[:14], tail))
    print("-" * 150)

    by_type = {}
    for r in rows:
        if r.get("scored") == "1" and not r.get("error"):
            by_type.setdefault(r["type"], []).append(r["hit_final"] == "1")
    print("\nHIT RATE BY TYPE")
    for t, hits in sorted(by_type.items()):
        print("  %-16s %d/%d" % (t, sum(hits), len(hits)))
    print()


def write_xlsx(paths, runs, summaries, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    F = "Arial"
    head = Font(name=F, size=10, bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1F4E79")
    body = Font(name=F, size=10)
    good = PatternFill("solid", fgColor="D6EFD8")
    bad = PatternFill("solid", fgColor="FBD5D5")
    warn = PatternFill("solid", fgColor="FFF0CC")

    wb = Workbook()

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append(["metric"] + [os.path.basename(p) for p in paths])
    for key, label, pattern in LABELS:
        ws.append([label] + [fmt(s[key], pattern) for s in summaries])
    for c in ws[1]:
        c.font = head
        c.fill = hfill
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = body
    ws.column_dimensions["A"].width = 32
    for i in range(len(paths)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 18
    ws.freeze_panes = "B2"

    # ---- Questions
    ws = wb.create_sheet("Questions")
    cols = ["id", "type", "question", "expected", "hit_final", "rank", "correct",
            "verdict", "answer", "evidence", "seconds", "prompt_tokens", "error"]
    ws.append(cols)
    for r in runs[0]:
        ws.append([r.get(c, "") for c in cols])
    for c in ws[1]:
        c.font = head
        c.fill = hfill
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = body
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (3, 9, 10))
        hit = row[4].value
        if hit == "1":
            row[4].fill = good
        elif hit == "0":
            row[4].fill = bad
        cor = row[6].value
        if cor == "1":
            row[6].fill = good
        elif cor == "0":
            row[6].fill = bad
        if row[12].value:
            row[12].fill = warn
    for col, width in zip("ABCDEFGHIJKLM",
                          [7, 14, 46, 16, 9, 7, 9, 12, 60, 34, 9, 13, 28]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C2"

    # ---- Scores
    ws = wb.create_sheet("Scores")
    cols = ["id", "vec_n", "vec_min", "vec_max", "vec_spread",
            "bm25_n", "bm25_min", "bm25_max", "bm25_spread", "bm25_zero_count",
            "hybrid_n", "hybrid_min", "hybrid_max", "hybrid_spread",
            "rerank_n", "rerank_raw_min", "rerank_raw_max", "rerank_raw_spread",
            "rerank_nan_count", "tie_fallback"]
    ws.append(cols)
    for r in runs[0]:
        ws.append([num(r.get(c)) if c not in ("id",) else r.get(c, "") for c in cols])
    for c in ws[1]:
        c.font = head
        c.fill = hfill
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = body
        if row[18].value:
            row[18].fill = bad          # rerank_nan_count
        if row[9].value:
            row[9].fill = warn          # bm25_zero_count
    ws.column_dimensions["A"].width = 7
    for i in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "B2"

    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    wb.save(out)
    print("workbook written to %s" % out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("csv", nargs="+", help="one or two result CSVs (before, after)")
    ap.add_argument("--xlsx", default=None, help="also write a formatted workbook")
    args = ap.parse_args()

    runs = [read(p) for p in args.csv]
    summaries = [summarize(r) for r in runs]
    print_terminal(args.csv, runs, summaries)
    if args.xlsx:
        write_xlsx(args.csv, runs, summaries, args.xlsx)


if __name__ == "__main__":
    main()